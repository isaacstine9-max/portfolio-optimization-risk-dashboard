"""
Moderna (NASDAQ: MRNA) discounted cash flow model.

Builds `Moderna_DCF.xlsx`, a fully formula-driven Excel workbook:
  * Inputs      - market data, WACC build, and Bear / Base / Bull scenario levers
  * Model       - revenue build (respiratory vaccines + risk-adjusted intismeran),
                  operating costs, NOL-shielded taxes, unlevered free cash flow
  * DCF         - discounting, terminal value, equity bridge, value per share,
                  and a WACC x terminal-growth sensitivity grid
  * Sources     - where every hard-coded number came from

All figures in US$ millions unless stated. Change the scenario selector
(Inputs!C6: 1 = Bear, 2 = Base, 3 = Bull) or any blue input and the whole
model recalculates in Excel.

Usage:
    python build_model.py                 # writes Moderna_DCF.xlsx
    python build_model.py -o out.xlsx     # custom output path
"""

import argparse
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# Styling
# --------------------------------------------------------------------------
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")            # hard-coded inputs
BLACK = Font(name=FONT, color="000000")           # formulas
GREEN = Font(name=FONT, color="008000")           # links to another sheet
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HEADER = Font(name=FONT, bold=True, color="FFFFFF")
NOTE = Font(name=FONT, italic=True, color="666666", size=9)

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
OUTPUT_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="999999")
TOP_BORDER = Border(top=THIN)

USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.00x'
FACTOR = '0.000'
YEAR = '0'

# --------------------------------------------------------------------------
# Data (all sourced - see SOURCES below)
# --------------------------------------------------------------------------
VALUATION_DATE = date(2026, 10, 5)
FIRST_YEAR, LAST_YEAR = 2027, 2040            # explicit forecast (valued)
YEARS = [2025, 2026] + list(range(FIRST_YEAR, LAST_YEAR + 1))

SRC_Q2 = ("Moderna Q2 2026 press release (8-K Ex. 99.1, Jul 31 2026): "
          "https://www.sec.gov/Archives/edgar/data/0001682852/000168285226000147/"
          "exhibit9912026q2pressrelea.htm")
SRC_FY25 = ("Moderna Q4/FY2025 press release (8-K Ex. 99.1, Feb 13 2026): "
            "https://www.sec.gov/Archives/edgar/data/1682852/000168285226000015/"
            "exhibit9912025q4pressrelea.htm")
SRC_JPM = ("Moderna JPM 2026 update (8-K Ex. 99.1, Jan 12 2026): "
           "https://www.sec.gov/Archives/edgar/data/1682852/000168285226000007/"
           "exhibit991-01122026.htm")
SRC_PRICE = ("Closing price Oct 2 2026 and market cap US$75.4bn: "
             "https://www.ad-hoc-news.de/boerse/news/corporate-news/"
             "moderna-stock-gained-0-57-percent-as-nasdaq-100-entry-nears/70226496")
SRC_UST = ("10-yr Treasury par yield 5.28% (Oct 2 2026): "
           "https://stockmarketwatch.com/bonds/10-year-treasury-yield")
SRC_WB = ("William Blair (Aug 2026) 2040 intismeran sales to Moderna after 50/50 "
          "Merck profit share: melanoma >$5.4bn, NSCLC ~$10bn, RCC ~$3.3bn: "
          "https://www.benzinga.com/analyst-stock-ratings/upgrades/26/08/61333408/"
          "modernas-intismeran-close-to-keytruda-style-oncology-opportunity-analyst")

# Scenario levers: (label, bear, base, bull, number format, note)
SCENARIOS = [
    ("Respiratory vaccines", None, None, None, None, None),
    ("Vaccine revenue growth 2027-2030", 0.02, 0.08, 0.15, PCT,
     "Flu (mFLUSIVA), flu+COVID (mCOMBRIAX), norovirus launches drive growth; "
     "company targets up to 10% growth in 2026."),
    ("Vaccine revenue growth 2031-2035", 0.00, 0.04, 0.06, PCT, "Assumption."),
    ("Vaccine revenue growth 2036-2040", -0.02, 0.02, 0.03, PCT, "Assumption."),
    ("Vaccine gross margin, steady state (from 2030)", 0.55, 0.65, 0.72, PCT,
     "2026E ~62% (guided ~$0.8bn cost of sales ex. $0.9bn settlement on ~$2.1bn "
     "revenue); management expects manufacturing improvements to lift margins."),
    ("Intismeran - adjuvant melanoma", None, None, None, None, None),
    ("Peak revenue to Moderna ($mm)", 2500, 4000, 5500, USD,
     "Moderna's half of the Merck 50/50 profit share. William Blair models >$5.4bn "
     "to Moderna by 2040."),
    ("Probability of approval", 0.80, 0.90, 0.95, PCT,
     "Phase 3 INTerpath-001 met RFS primary endpoint (Aug 19 2026)."),
    ("Intismeran - NSCLC", None, None, None, None, None),
    ("Peak revenue to Moderna ($mm)", 2000, 5000, 10000, USD,
     "William Blair models ~$10bn to Moderna by 2040 (bull case)."),
    ("Probability of approval", 0.35, 0.50, 0.60, PCT, "Phase 3 ongoing; assumption."),
    ("Intismeran - RCC, bladder & other", None, None, None, None, None),
    ("Peak revenue to Moderna ($mm)", 500, 1500, 3300, USD,
     "William Blair models ~$3.3bn RCC to Moderna by 2040 (bull case)."),
    ("Probability of approval", 0.30, 0.45, 0.55, PCT, "Phase 2/3 ongoing; assumption."),
    ("Intismeran economics & terminal value", None, None, None, None, None),
    ("Intismeran cost of sales (% of Moderna revenue)", 0.35, 0.25, 0.20, PCT,
     "Individualized manufacturing is costlier than vaccines; assumption."),
    ("Terminal growth rate", 0.015, 0.025, 0.030, PCT, "Applied after 2040."),
    ("Cost base", None, None, None, None, None),
    ("R&D floor ($mm, from 2028)", 1600, 2400, 2800, USD,
     "Bear: management has said it would cut costs further if sales miss. "
     "Bull: more oncology investment."),
    ("SG&A floor ($mm, from 2028)", 700, 900, 1000, USD,
     "2026E guidance ~$1.0bn; continued cuts."),
]

# Common (non-scenario) operating assumptions: (label, value, fmt, note)
OPERATING = [
    ("Launch year - melanoma", 2028, YEAR, "Assumes filing after Phase 3 interim; approval in 2027-28."),
    ("Launch year - NSCLC", 2030, YEAR, "Assumption."),
    ("Launch year - RCC, bladder & other", 2030, YEAR, "Assumption."),
    ("Years from launch to peak", 6, NUM, "Linear ramp; assumption."),
    ("Vaccine gross margin 2026E", 0.62, PCT, SRC_Q2),
    ("R&D as % of revenue (if above floor)", 0.20, PCT, "Large-cap biopharma norm ~18-22%."),
    ("SG&A as % of revenue (if above floor)", 0.12, PCT, "Merck co-commercializes intismeran; assumption."),
    ("D&A floor ($mm)", 350, USD, "GAAP opex less cash costs ~ $0.7bn incl. SBC; assumption."),
    ("D&A as % of revenue", 0.03, PCT, "Assumption."),
    ("Capex floor ($mm)", 250, USD, "2026 guidance $0.2-0.3bn capex."),
    ("Capex as % of revenue", 0.035, PCT, "Assumption."),
    ("Net working capital, % of revenue change", 0.15, PCT, "Assumption."),
    ("Long-run cash tax rate", 0.21, PCT, "US federal statutory; assumption."),
    ("Usable NOLs at YE2026 ($mm)", 12000, USD,
     "Cumulative 2023-2026 losses ~US$14bn; usable amount is an estimate."),
    ("NOL usage cap (% of taxable income)", 0.80, PCT, "Post-2017 US federal NOL rule."),
]


def style_header_row(ws, row, first_col, last_col):
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def section(ws, row, text, last_col):
    ws.cell(row=row, column=1, value=text).font = BOLD
    for c in range(1, last_col + 1):
        ws.cell(row=row, column=c).fill = SECTION_FILL


def put(ws, ref, value, font=BLACK, fmt=None, fill=None, note=None):
    cell = ws[ref]
    cell.value = value
    cell.font = font
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if note:
        cell.comment = Comment(note, "Model")
    return cell


# --------------------------------------------------------------------------
# Inputs sheet
# --------------------------------------------------------------------------
def build_inputs(wb):
    ws = wb.active
    ws.title = "Inputs"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 50
    for col in "CDEF":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["G"].width = 90

    ws["B1"] = "Moderna (MRNA) DCF - Inputs"
    ws["B1"].font = TITLE
    ws["B2"] = ("US$ millions unless stated. Blue = hard-coded input, black = formula, "
                "green = link from another sheet. Yellow = key lever. "
                "Hover over cells with a red corner for sources.")
    ws["B2"].font = NOTE

    refs = {}

    # Scenario selector
    put(ws, "B4", "Scenario selector", BOLD)
    put(ws, "B6", "Active scenario (1 = Bear, 2 = Base, 3 = Bull)")
    put(ws, "C6", 2, BLUE, NUM, KEY_FILL,
        "Change to 1, 2 or 3; the Live column and every sheet update.")
    put(ws, "D6", '=CHOOSE(C6,"Bear","Base","Bull")', BLACK)
    refs["sel"] = "Inputs!$C$6"
    refs["scen_name"] = "Inputs!$D$6"

    # Market data
    r = 8
    section(ws, r, "Market data & balance sheet", 7)
    rows = [
        ("valdate", "Valuation date", VALUATION_DATE, "yyyy-mm-dd", "Model date."),
        ("price", "Share price (US$)", 190.01, USD2, SRC_PRICE),
        ("shares", "Diluted shares outstanding (mm)", None, NUM1,
         "Market cap US$75.4bn / price US$190.01. " + SRC_PRICE),
        ("cash", "Cash & investments, YE2026E ($mm)", 4950, USD,
         "Midpoint of US$4.7-5.2bn year-end 2026 guidance (already net of the "
         "$950mm July 2026 settlement and remaining 2026 burn). " + SRC_Q2),
        ("debt", "Debt ($mm)", 600, USD,
         "US$0.6bn drawn on the US$1.5bn Ares term loan; balance sheet long-term "
         "debt US$591mm at Jun 30 2026. " + SRC_JPM),
    ]
    r += 1
    for key, label, val, fmt, note in rows:
        put(ws, f"B{r}", label)
        if key == "shares":
            put(ws, f"C{r}", "=75400/C10", BLACK, fmt, note=note)
        else:
            put(ws, f"C{r}", val, BLUE, fmt, note=note)
        put(ws, f"G{r}", note, NOTE)
        refs[key] = f"Inputs!$C${r}"
        r += 1
    put(ws, f"B{r}", "Net cash ($mm)", BOLD)
    put(ws, f"C{r}", "=C12-C13", BLACK, USD)
    refs["netcash"] = f"Inputs!$C${r}"
    put(ws, f"B{r+1}", "Market capitalization ($mm)")
    put(ws, f"C{r+1}", "=C10*C11", BLACK, USD)
    refs["mktcap"] = f"Inputs!$C${r+1}"
    r += 3

    # WACC
    section(ws, r, "Discount rate (WACC)", 7)
    r += 1
    wacc_rows = [
        ("rf", "Risk-free rate (10-yr UST)", 0.0528, PCT, SRC_UST),
        ("beta", "Equity beta", 1.30, "0.00",
         "Assumption; MRNA's historical beta has ranged ~1.2-1.8."),
        ("erp", "Equity risk premium", 0.050, PCT, "Assumption (Damodaran-style implied ERP ~4.5-5.5%)."),
        ("kd", "Pre-tax cost of debt", 0.090, PCT, "Approx. rate on Ares term loan; assumption."),
    ]
    start = r
    for key, label, val, fmt, note in wacc_rows:
        put(ws, f"B{r}", label)
        put(ws, f"C{r}", val, BLUE, fmt, note=note)
        put(ws, f"G{r}", note, NOTE)
        refs[key] = f"Inputs!$C${r}"
        r += 1
    rf, beta, erp, kd = (f"C{start}", f"C{start+1}", f"C{start+2}", f"C{start+3}")
    put(ws, f"B{r}", "Cost of equity (CAPM)")
    put(ws, f"C{r}", f"={rf}+{beta}*{erp}", BLACK, PCT)
    ke = f"C{r}"
    r += 1
    put(ws, f"B{r}", "After-tax cost of debt")
    tax_row_placeholder = r  # filled after operating block is placed
    kd_at = f"C{r}"
    r += 1
    put(ws, f"B{r}", "Equity weight")
    mc, dt = refs["mktcap"].split("!")[1], refs["debt"].split("!")[1]
    put(ws, f"C{r}", f"={mc}/({mc}+{dt})", BLACK, PCT)
    we = f"C{r}"
    r += 1
    put(ws, f"B{r}", "Debt weight")
    put(ws, f"C{r}", f"=1-{we}", BLACK, PCT)
    wd = f"C{r}"
    r += 1
    put(ws, f"B{r}", "WACC (calculated)", BOLD)
    put(ws, f"C{r}", f"={we}*{ke}+{wd}*{kd_at}", BLACK, PCT)
    wacc_calc = f"C{r}"
    r += 1
    put(ws, f"B{r}", "WACC override (leave blank to use calculated)")
    put(ws, f"C{r}", None, BLUE, PCT, KEY_FILL, "Type a rate here to override the CAPM build.")
    override = f"C{r}"
    r += 1
    put(ws, f"B{r}", "WACC used", BOLD)
    put(ws, f"C{r}", f'=IF(ISNUMBER({override}),{override},{wacc_calc})', BLACK, PCT, OUTPUT_FILL)
    refs["wacc"] = f"Inputs!$C${r}"
    r += 2

    # Scenario table
    section(ws, r, "Scenario levers", 7)
    r += 1
    for c, h in zip("BCDEFG", ["Lever", "Bear", "Base", "Bull", "Live", "Rationale / source"]):
        ws[f"{c}{r}"] = h
    style_header_row(ws, r, 2, 7)
    r += 1
    scen_keys = ["", "g1", "g2", "g3", "vgm", "", "mel_peak", "mel_pos", "",
                 "nsclc_peak", "nsclc_pos", "", "oth_peak", "oth_pos", "",
                 "onc_cogs", "tg", "", "rd_floor", "sga_floor"]
    for (label, bear, base, bull, fmt, note), key in zip(SCENARIOS, scen_keys):
        if bear is None:
            ws[f"B{r}"] = label
            ws[f"B{r}"].font = BOLD
            r += 1
            continue
        put(ws, f"B{r}", "   " + label)
        for col, v in zip("CDE", (bear, base, bull)):
            put(ws, f"{col}{r}", v, BLUE, fmt)
        put(ws, f"F{r}", f"=CHOOSE({refs['sel'].split('!')[1]},C{r},D{r},E{r})",
            BLACK, fmt, KEY_FILL)
        put(ws, f"G{r}", note, NOTE)
        refs[key] = f"Inputs!$F${r}"
        r += 1
    r += 1

    # Operating assumptions
    section(ws, r, "Operating assumptions (all scenarios)", 7)
    r += 1
    op_keys = ["mel_launch", "nsclc_launch", "oth_launch", "ytp", "vgm26",
               "rd_pct", "sga_pct", "da_floor", "da_pct",
               "capex_floor", "capex_pct", "nwc_pct", "tax", "nol", "nol_cap"]
    for (label, val, fmt, note), key in zip(OPERATING, op_keys):
        put(ws, f"B{r}", label)
        put(ws, f"C{r}", val, BLUE, fmt, note=note)
        put(ws, f"G{r}", note, NOTE)
        refs[key] = f"Inputs!$C${r}"
        r += 1

    # After-tax cost of debt now that the tax row exists
    tax_local = refs["tax"].split("!")[1]
    put(ws, kd_at, f"={kd}*(1-{tax_local})", BLACK, PCT)
    ws.freeze_panes = "C4"
    return refs


# --------------------------------------------------------------------------
# Model sheet
# --------------------------------------------------------------------------
def build_model(wb, R):
    ws = wb.create_sheet("Model")
    ws.column_dimensions["A"].width = 46
    ncols = len(YEARS)
    for i in range(ncols):
        ws.column_dimensions[get_column_letter(i + 2)].width = 11

    ws["A1"] = "Moderna (MRNA) - Operating model & free cash flow (US$ mm)"
    ws["A1"].font = TITLE
    ws["A2"] = '="Scenario: "&' + R["scen_name"]
    ws["A2"].font = GREEN

    col = {y: get_column_letter(i + 2) for i, y in enumerate(YEARS)}
    last = ncols + 1

    # Header rows
    ws["A4"] = "Fiscal year (Dec)"
    for y in YEARS:
        put(ws, f"{col[y]}4", y, BLUE if y <= 2026 else BLACK, YEAR)
        ws[f"{col[y]}5"] = "Actual" if y == 2025 else ("Guide" if y == 2026 else "Forecast")
    style_header_row(ws, 4, 1, last)
    for c in range(1, last + 1):
        ws.cell(row=5, column=c).font = NOTE
        ws.cell(row=5, column=c).alignment = Alignment(horizontal="center")
    for y in YEARS[2:]:
        put(ws, f"{col[y]}4", f"={col[y-1]}4+1", HEADER, YEAR)

    def row_label(r, text, bold=False):
        ws[f"A{r}"] = text
        ws[f"A{r}"].font = BOLD if bold else BLACK

    def fill_row(r, fn, fmt=USD, years=None, font=BLACK, top=False):
        for y in (years or YEARS):
            c = col[y]
            put(ws, f"{c}{r}", fn(y, c), font, fmt)
            if top:
                ws[f"{c}{r}"].border = TOP_BORDER

    fy = YEARS[2:]          # forecast years
    p = lambda y: col[y - 1]  # previous column letter

    # ---------------- Revenue ----------------
    section(ws, 7, "Revenue", last)
    row_label(8, "Respiratory vaccines (COVID, RSV, flu, combo, noro)")
    put(ws, "B8", 1944, BLUE, USD, note="FY2025 total revenue US$1.9bn. " + SRC_FY25)
    put(ws, "C8", "=B8*1.08", BLACK, USD,
        note="2026 guide: up to 10% growth; modeled at +8% (consensus ~+8%).")
    fill_row(8, lambda y, c:
             f"={p(y)}8*(1+IF({c}$4<=2030,{R['g1']},IF({c}$4<=2035,{R['g2']},{R['g3']})))",
             years=fy)

    def ramp(launch, peak, pos):
        return lambda y, c: (f"=IF({c}$4>={launch},MIN(1,({c}$4-{launch}+1)/{R['ytp']}),0)"
                             f"*{peak}*{pos}")

    row_label(9, "Intismeran - melanoma (risk-adjusted, Moderna share)")
    fill_row(9, ramp(R["mel_launch"], R["mel_peak"], R["mel_pos"]), years=fy)
    row_label(10, "Intismeran - NSCLC (risk-adjusted)")
    fill_row(10, ramp(R["nsclc_launch"], R["nsclc_peak"], R["nsclc_pos"]), years=fy)
    row_label(11, "Intismeran - RCC, bladder & other (risk-adjusted)")
    fill_row(11, ramp(R["oth_launch"], R["oth_peak"], R["oth_pos"]), years=fy)
    for y in YEARS[:2]:
        for rr in (9, 10, 11):
            put(ws, f"{col[y]}{rr}", 0, BLUE, USD)
    row_label(12, "Total intismeran")
    fill_row(12, lambda y, c: f"=SUM({c}9:{c}11)")
    row_label(13, "Total revenue", True)
    fill_row(13, lambda y, c: f"={c}8+{c}12", top=True)
    for y in YEARS:
        ws[f"{col[y]}13"].font = BOLD
    row_label(14, "   growth %")
    fill_row(14, lambda y, c: f"=IF({p(y)}13=0,0,{c}13/{p(y)}13-1)", PCT, years=YEARS[1:])

    # ---------------- Costs ----------------
    section(ws, 16, "Operating costs", last)
    row_label(17, "Vaccine gross margin %")
    put(ws, "C17", "=" + R["vgm26"], GREEN, PCT)
    fill_row(17, lambda y, c:
             f"=IF({c}$4>=2030,{R['vgm']},{R['vgm26']}+({R['vgm']}-{R['vgm26']})*({c}$4-2026)/4)",
             PCT, years=fy)
    row_label(18, "Cost of sales - vaccines")
    fill_row(18, lambda y, c: f"={c}8*(1-{c}17)", years=YEARS[1:])
    ws["C18"].comment = Comment("2026 guide US$1.7bn incl. US$0.9bn one-time settlement; "
                                "modeled ex-settlement (~US$0.8bn). " + SRC_Q2, "Model")
    row_label(19, "Cost of sales - intismeran")
    fill_row(19, lambda y, c: f"={c}12*{R['onc_cogs']}", years=YEARS[1:])
    row_label(20, "Gross profit", True)
    fill_row(20, lambda y, c: f"={c}13-{c}18-{c}19", years=YEARS[1:], top=True)
    row_label(21, "   gross margin %")
    fill_row(21, lambda y, c: f"=IF({c}13=0,0,{c}20/{c}13)", PCT, years=YEARS[1:])

    row_label(22, "Research & development")
    put(ws, "C22", 2900, BLUE, USD, note="2026 guide ~US$2.9bn. " + SRC_Q2)
    put(ws, "D22", 2400, BLUE, USD,
        note="2027 GAAP opex guided US$4.2-4.6bn, later improved ~US$0.5bn; "
             "R&D implied ~US$2.3-2.5bn. " + SRC_JPM)
    fill_row(22, lambda y, c: f"=MAX({R['rd_floor']},{R['rd_pct']}*{c}13)",
             years=list(range(2028, LAST_YEAR + 1)))
    row_label(23, "Selling, general & administrative")
    put(ws, "C23", 1000, BLUE, USD, note="2026 guide ~US$1.0bn. " + SRC_Q2)
    put(ws, "D23", 950, BLUE, USD, note="Assumption: continued cost cuts.")
    fill_row(23, lambda y, c: f"=MAX({R['sga_floor']},{R['sga_pct']}*{c}13)",
             years=list(range(2028, LAST_YEAR + 1)))
    row_label(24, "Operating income (EBIT)", True)
    fill_row(24, lambda y, c: f"={c}20-{c}22-{c}23", years=YEARS[1:], top=True)
    for y in YEARS[1:]:
        ws[f"{col[y]}24"].font = BOLD
    row_label(25, "   EBIT margin %")
    fill_row(25, lambda y, c: f"=IF({c}13=0,0,{c}24/{c}13)", PCT, years=YEARS[1:])

    # ---------------- Taxes with NOLs ----------------
    section(ws, 27, "Cash taxes (NOL shield)", last)
    row_label(28, "NOL balance - beginning")
    put(ws, f"{col[FIRST_YEAR]}28", "=" + R["nol"], GREEN, USD)
    fill_row(28, lambda y, c: f"={p(y)}30", years=fy[1:])
    row_label(29, "NOL utilized")
    fill_row(29, lambda y, c: f"=MIN({c}28,{R['nol_cap']}*MAX({c}24,0))", years=fy)
    row_label(30, "NOL balance - ending")
    fill_row(30, lambda y, c: f"={c}28-{c}29", years=fy)
    row_label(31, "Taxable income")
    fill_row(31, lambda y, c: f"=MAX({c}24,0)-{c}29", years=fy)
    row_label(32, "Cash taxes")
    fill_row(32, lambda y, c: f"={c}31*{R['tax']}", years=fy)

    # ---------------- Free cash flow ----------------
    section(ws, 34, "Unlevered free cash flow", last)
    row_label(35, "EBIT")
    fill_row(35, lambda y, c: f"={c}24", years=fy)
    row_label(36, "less: cash taxes")
    fill_row(36, lambda y, c: f"=-{c}32", years=fy)
    row_label(37, "plus: D&A")
    fill_row(37, lambda y, c: f"=MAX({R['da_floor']},{R['da_pct']}*{c}13)", years=fy)
    row_label(38, "less: capital expenditures")
    fill_row(38, lambda y, c: f"=-MAX({R['capex_floor']},{R['capex_pct']}*{c}13)", years=fy)
    row_label(39, "less: increase in net working capital")
    fill_row(39, lambda y, c: f"=-{R['nwc_pct']}*({c}13-{p(y)}13)", years=fy)
    row_label(40, "Unlevered free cash flow", True)
    fill_row(40, lambda y, c: f"=SUM({c}35:{c}39)", years=fy, top=True)
    for y in fy:
        ws[f"{col[y]}40"].font = BOLD
        ws[f"{col[y]}40"].fill = OUTPUT_FILL

    # ---------------- Discounting ----------------
    section(ws, 42, "Discounting (mid-year convention, from valuation date)", last)
    row_label(43, "Discount period (years)")
    fill_row(43, lambda y, c:
             f"=(DATE({c}$4,6,30)-{R['valdate']})/365", FACTOR, years=fy)
    row_label(44, "Discount factor")
    fill_row(44, lambda y, c: f"=1/(1+{R['wacc']})^{c}43", FACTOR, years=fy)
    row_label(45, "PV of free cash flow", True)
    fill_row(45, lambda y, c: f"={c}40*{c}44", years=fy)

    ws.freeze_panes = "B6"
    first, lastc = col[FIRST_YEAR], col[LAST_YEAR]
    return {
        "fcf": f"Model!${first}$40:${lastc}$40",
        "t": f"Model!${first}$43:${lastc}$43",
        "pv": f"Model!${first}$45:${lastc}$45",
        "fcf_last": f"Model!${lastc}$40",
        "year_last": f"Model!${lastc}$4",
        "rev_last": f"Model!${lastc}$13",
        "ebit_margin_last": f"Model!${lastc}$25",
        "intis_last": f"Model!${lastc}$12",
    }


# --------------------------------------------------------------------------
# DCF sheet
# --------------------------------------------------------------------------
def build_dcf(wb, R, M):
    ws = wb.create_sheet("DCF")
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 44
    for c in "CDEFGHI":
        ws.column_dimensions[c].width = 13

    ws["B1"] = "Moderna (MRNA) - DCF valuation"
    ws["B1"].font = TITLE
    put(ws, "B2", '="Scenario: "&' + R["scen_name"] + '&"   |   WACC "&TEXT(' + R["wacc"]
        + ',"0.0%")&"   |   terminal growth "&TEXT(' + R["tg"] + ',"0.0%")', GREEN)

    rows = [
        ("Sum of PV of free cash flows, 2027-2040", f"=SUM({M['pv']})", USD),
        ("Terminal-year FCF x (1 + g)", f"={M['fcf_last']}*(1+{R['tg']})", USD),
        ("Terminal value at end-2040", f"=C5/({R['wacc']}-{R['tg']})", USD),
        ("Discount period to end-2040 (years)",
         f"=(DATE({M['year_last']},12,31)-{R['valdate']})/365", FACTOR),
        ("PV of terminal value", f"=C6/(1+{R['wacc']})^C7", USD),
        ("Enterprise value", "=C4+C8", USD),
        ("plus: net cash (YE2026E cash less debt)", "=" + R["netcash"], USD),
        ("Equity value", "=C9+C10", USD),
        ("Diluted shares (mm)", "=" + R["shares"], NUM1),
        ("Implied value per share (US$)", "=C11/C12", USD2),
        ("Current share price (US$)", "=" + R["price"], USD2),
        ("Upside / (downside)", "=C13/C14-1", PCT),
    ]
    section(ws, 3, "Valuation summary (US$ mm)", 9)
    for i, (label, f, fmt) in enumerate(rows):
        r = 4 + i
        put(ws, f"B{r}", label)
        font = GREEN if f.startswith("=Inputs") or f.startswith("=SUM(Model") else BLACK
        put(ws, f"C{r}", f, font, fmt)
    for r in (9, 11, 13, 15):
        ws[f"B{r}"].font = BOLD
        ws[f"C{r}"].font = BOLD
        ws[f"C{r}"].fill = OUTPUT_FILL
    ws["C13"].fill = KEY_FILL

    section(ws, 17, "Diagnostics", 9)
    diag = [
        ("Terminal value as % of enterprise value", "=IF(C9=0,0,C8/C9)", PCT),
        ("Implied EV / 2040E revenue", f"=IF({M['rev_last']}=0,0,C9/{M['rev_last']})", MULT),
        ("2040E revenue ($mm)", "=" + M["rev_last"], USD),
        ("2040E risk-adjusted intismeran revenue ($mm)", "=" + M["intis_last"], USD),
        ("2040E EBIT margin", "=" + M["ebit_margin_last"], PCT),
        ("Market cap at current price ($mm)", "=" + R["mktcap"], USD),
        ("Market-implied enterprise value ($mm)", "=C23-" + R["netcash"], USD),
    ]
    for i, (label, f, fmt) in enumerate(diag):
        r = 18 + i
        put(ws, f"B{r}", label)
        put(ws, f"C{r}", f, GREEN if "Model!" in f or "Inputs!" in f else BLACK, fmt)

    # Sensitivity: value per share vs WACC (rows) and terminal growth (cols)
    top = 27
    section(ws, top, "Sensitivity - implied value per share (US$): WACC (rows) vs terminal growth (columns)", 9)
    put(ws, f"B{top+1}", "WACC step / growth step")
    put(ws, f"C{top+1}", 0.01, BLUE, PCT, note="Row spacing for WACC.")
    put(ws, f"D{top+1}", 0.005, BLUE, PCT, note="Column spacing for terminal growth.")
    hdr = top + 2
    put(ws, f"B{hdr}", "WACC \\ g", BOLD)
    gcols = "CDEFG"
    for j, c in enumerate(gcols):
        put(ws, f"{c}{hdr}", f"={R['tg']}+({j}-2)*$D${top+1}", BOLD, PCT)
    for i in range(7):
        r = hdr + 1 + i
        put(ws, f"B{r}", f"={R['wacc']}+({i}-3)*$C${top+1}", BOLD, PCT)
        for c in gcols:
            w, g = f"$B{r}", f"{c}${hdr}"
            f = (f"=(SUMPRODUCT({M['fcf']},1/(1+{w})^{M['t']})"
                 f"+{M['fcf_last']}*(1+{g})/({w}-{g})/(1+{w})^$C$7"
                 f"+{R['netcash']})/{R['shares']}")
            put(ws, f"{c}{r}", f, BLACK, USD2)
            if i == 3 and c == "E":
                ws[f"{c}{r}"].fill = KEY_FILL
    put(ws, f"B{hdr+9}", "Centre cell matches the live valuation above. "
        "Rows and columns are re-centred on the live WACC and growth rate.", NOTE)
    ws.freeze_panes = "A4"


# --------------------------------------------------------------------------
# Sources sheet
# --------------------------------------------------------------------------
def build_sources(wb):
    ws = wb.create_sheet("Sources")
    ws.column_dimensions["A"].width = 160
    ws["A1"] = "Sources and notes"
    ws["A1"].font = TITLE
    lines = [
        SRC_FY25, SRC_Q2, SRC_JPM, SRC_PRICE, SRC_UST, SRC_WB,
        "",
        "Method notes:",
        "- Valued on unlevered FCF for 2027-2040; 2026 is shown for reference only because "
        "its remaining cash burn is already reflected in the YE2026 cash guidance used in "
        "the equity bridge.",
        "- Intismeran revenue = Moderna's share under the 50/50 Merck collaboration, ramped "
        "linearly to peak and multiplied by probability of approval (risk-adjusted).",
        "- Stock-based compensation is treated as a real cost (left inside opex, not added back).",
        "- Interest income on cash is excluded from FCF; cash enters through the equity bridge.",
        "- Not investment advice. All forecasts and probabilities are illustrative assumptions.",
    ]
    for i, line in enumerate(lines, start=3):
        ws[f"A{i}"] = line
        ws[f"A{i}"].font = BOLD if line.endswith(":") else Font(name=FONT)


def build(path):
    wb = Workbook()
    R = build_inputs(wb)
    M = build_model(wb, R)
    build_dcf(wb, R, M)
    build_sources(wb)
    wb.move_sheet("DCF", offset=-2)          # DCF first: Inputs, Model order after
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-o", "--output", default="Moderna_DCF.xlsx")
    args = ap.parse_args()
    print("Wrote", build(args.output))
